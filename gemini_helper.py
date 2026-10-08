"""All Gemini (Google AI Studio free API key) calls live here.

Uses the plain REST API through `requests`, so there is no SDK version to break.
The key is read from Streamlit secrets / environment variable - never hard-coded.
"""
import base64
import io
import json
import os
import re
import time

import requests
import streamlit as st

import scoring as sc

API_BASE = "https://generativelanguage.googleapis.com/v1beta/models"
# Tried in this order. Gemini 2.5 models are scheduled for shutdown in Oct 2026, so newer
# models / the "latest" alias come first. Override with GEMINI_MODEL in secrets if you like.
DEFAULT_MODELS = ["gemini-flash-latest", "gemini-flash-lite-latest", "gemini-3-flash-preview", "gemini-2.5-flash"]
_state = {"model": None}

COMPANY = "M Beans & Bites Pvt. Ltd."


class GeminiError(Exception):
    """Raised for any Gemini failure; the message is safe to show to the user."""


# ------------------------------------------------------------ configuration
def _secret(name):
    try:
        if name in st.secrets:
            return str(st.secrets[name]).strip()
    except Exception:
        pass
    return (os.environ.get(name) or "").strip()


def get_api_key():
    k = st.session_state.get("gemini_api_key")
    if k:
        return k.strip()
    return _secret("GEMINI_API_KEY") or _secret("GOOGLE_API_KEY") or None


def _candidate_models():
    models = []
    custom = _secret("GEMINI_MODEL")
    if custom:
        models.append(custom)
    if _state["model"]:
        models.append(_state["model"])
    models += DEFAULT_MODELS
    seen, out = set(), []
    for m in models:
        if m not in seen:
            seen.add(m)
            out.append(m)
    return out


def last_model():
    return _state["model"]


# ----------------------------------------------------------------- low level
def _err_message(resp):
    try:
        return resp.json().get("error", {}).get("message", resp.text[:200])
    except Exception:
        return resp.text[:200]


def _extract_text(data):
    fb = data.get("promptFeedback") or {}
    if fb.get("blockReason"):
        raise GeminiError(f"Gemini blocked the request ({fb['blockReason']}).")
    cands = data.get("candidates") or []
    if not cands:
        return ""
    parts = (cands[0].get("content") or {}).get("parts") or []
    return "".join(p.get("text", "") for p in parts if p.get("text") and not p.get("thought")).strip()


def generate(parts, system=None, json_mode=False, max_tokens=4096):
    """Call Gemini and return the text. 'Thinking' is switched off for speed.
    On rate-limit / model errors it moves straight to the next model (no long waits)."""
    key = get_api_key()
    if not key:
        raise GeminiError("No Gemini API key found. Add GEMINI_API_KEY in the Streamlit secrets "
                          "(or paste a key in the sidebar).")
    base_cfg = {"maxOutputTokens": max_tokens, "temperature": 0.2}
    if json_mode:
        base_cfg["responseMimeType"] = "application/json"
    last_err = "Unknown error"
    for model in _candidate_models():
        use_think_cfg, tries = True, 0
        while tries < 3:
            tries += 1
            cfg = dict(base_cfg)
            if use_think_cfg:
                cfg["thinkingConfig"] = {"thinkingBudget": 0}
            body = {"contents": [{"role": "user", "parts": parts}], "generationConfig": cfg}
            if system:
                body["systemInstruction"] = {"parts": [{"text": system}]}
            try:
                resp = requests.post(f"{API_BASE}/{model}:generateContent",
                                     headers={"x-goog-api-key": key, "Content-Type": "application/json"},
                                     json=body, timeout=60)
            except requests.RequestException as e:
                last_err = f"Network problem while contacting Gemini ({type(e).__name__})."
                break
            code = resp.status_code
            if code == 200:
                text = _extract_text(resp.json())
                if text:
                    _state["model"] = model
                    return text
                last_err = "Gemini returned an empty response."
                break
            msg = _err_message(resp)
            if code == 400 and use_think_cfg and "think" in msg.lower():
                use_think_cfg = False          # model does not accept the setting -> retry without it
                continue
            if code in (401, 403) or "API key" in msg or "API_KEY" in msg:
                raise GeminiError("Gemini rejected the API key. Please check GEMINI_API_KEY. "
                                  f"(Details: {msg[:150]})")
            if code == 429:
                last_err = "Gemini free-tier limit reached. Please wait a minute and try again."
                break                          # try the next model (separate quota)
            if code >= 500:
                last_err = f"Gemini service error ({code})."
                time.sleep(1)
                continue
            last_err = f"Gemini error {code}: {msg[:200]}"
            break                              # e.g. model not found -> next model
    raise GeminiError(last_err)


def _parse_json(text):
    t = text.strip()
    t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.I)
    t = re.sub(r"\s*```$", "", t)
    try:
        return json.loads(t)
    except Exception:
        m = re.search(r"(\{.*\}|\[.*\])", t, re.S)
        if m:
            try:
                return json.loads(m.group(1))
            except Exception:
                pass
    raise GeminiError("Gemini returned a response that could not be read (invalid JSON).")


def generate_json(parts, system=None):
    last = None
    for _ in range(2):
        try:
            return _parse_json(generate(parts, system=system, json_mode=True))
        except GeminiError as e:
            last = e
            if "invalid JSON" not in str(e):
                break
    raise last


def test_connection():
    text = generate([{"text": "Reply with the single word: OK"}])
    return f"Connected (model: {_state['model']})", text


# ------------------------------------------------------- document handling
IMAGE_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp"}
SUPPORTED_TYPES = ["pdf", "png", "jpg", "jpeg", "webp", "docx", "xlsx", "csv", "txt"]
MAX_INLINE_BYTES = 18 * 1024 * 1024


def _docx_text(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _xlsx_text(data):
    import pandas as pd
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    out = []
    for name, df in sheets.items():
        out.append(f"[Sheet: {name}]")
        out.append(df.fillna("").to_csv(index=False, header=False))
    return "\n".join(out)


def prepare_documents(files):
    """Turn uploaded files into Gemini 'parts'. Returns (parts, ok_names, failed_names)."""
    parts, ok, failed, total = [], [], [], 0
    for f in files:
        name = f.name
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        try:
            data = f.getvalue()
            if ext == "pdf" or ext in IMAGE_MIME:
                total += len(data)
                if total > MAX_INLINE_BYTES or not data:
                    raise ValueError("size")
                mime = "application/pdf" if ext == "pdf" else IMAGE_MIME[ext]
                parts.append({"text": f"=== DOCUMENT: {name} ==="})
                parts.append({"inline_data": {"mime_type": mime, "data": base64.b64encode(data).decode()}})
            else:
                if ext == "docx":
                    text = _docx_text(data)
                elif ext == "xlsx":
                    text = _xlsx_text(data)
                elif ext in ("csv", "txt"):
                    text = data.decode("utf-8", errors="replace")
                else:
                    raise ValueError("unsupported")
                if not text.strip():
                    raise ValueError("empty")
                parts.append({"text": f"=== DOCUMENT: {name} ===\n{text[:60000]}"})
            ok.append(name)
        except Exception:
            failed.append(name)
    return parts, ok, failed


# ---------------------------------------------------------- 1. extraction
EXTRACT_SYSTEM = """You are a document-understanding engine for the procurement team of M Beans & Bites, a café business in India.
You receive one or more vendor documents (GST certificate, bank details / cancelled cheque, quotation, company profile)
plus a list of EXISTING vendors. Return ONLY one JSON object - no commentary.

STRICT RULES
- NEVER invent or guess. If a value is not clearly present in the documents use null.
- Copy identifiers (GSTIN, PAN, account number, IFSC) exactly as printed, without spaces.
- unit_price = price in rupees for ONE kg (coffee beans) or ONE cup (paper cups). If quoted per pack/carton/1000 units,
  convert by simple division and mention it in other_terms. If conversion is impossible use null.
- Include EVERY product in the quotation, even if it is not coffee beans or paper cups.
- delivery_days = whole number of days (null if not stated).
- quality_evidence: copy ONLY what the documents state (certification names, on-time delivery %, defect/return %).
- inconsistencies: real mismatches between documents (vendor name vs account holder name, PAN vs PAN inside GSTIN,
  different addresses). Empty list if none.
- duplicate_matches: compare the vendor's legal name and address with the EXISTING vendors (spelling variants,
  abbreviations, 'Supply' vs 'Suppliers', 'Estate' vs 'Area'). Include only those with similarity >= 60. Empty list if none.
- summary: 2 short business sentences (max 50 words) about who the vendor is, what it quotes (price, delivery) and
  whether the documentation looks complete. Facts from the documents only.

JSON SHAPE
{
 "documents": [{"file_name": "string", "document_type": "GST Certificate | Bank Details | Quotation | Company Profile | Other"}],
 "vendor": {"legal_name": null, "trade_name": null, "gstin": null, "pan": null, "address": null,
            "contact_person": null, "email": null, "phone": null, "category": null, "products_supplied": []},
 "bank": {"bank_name": null, "account_holder": null, "account_number": null, "ifsc": null},
 "quotations": [{"product_name": "string", "product_code": null, "unit_price": null, "quoted_quantity": null,
                 "total_price": null, "taxes": null, "delivery_days": null, "validity": null, "other_terms": null}],
 "quality_evidence": {"certifications": [], "on_time_delivery_percent": null, "defect_or_return_rate_percent": null,
                      "other_evidence": null},
 "inconsistencies": [],
 "duplicate_matches": [{"vendor_id": "V001", "similarity": 0, "reasons": ["Similar vendor name"]}],
 "summary": "string"
}"""


def to_number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    if s in ("", "-", ".", "-."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _s(v):
    s = sc.clean(v)
    return s or None


def extract_vendor_documents(parts, existing=None):
    """Run Gemini document understanding. Returns a validated, normalised dict."""
    if not parts:
        raise GeminiError("No readable documents were provided.")
    existing = existing or []
    ex_list = [{"vendor_id": v["vendor_id"], "legal_name": v["legal_name"], "address": v["address"]} for v in existing]
    prompt = parts + [{"text": "EXISTING vendors (JSON): " + json.dumps(ex_list, ensure_ascii=False) +
                               "\n\nExtract the vendor information from the documents above and return the JSON."}]
    raw = generate_json(prompt, system=EXTRACT_SYSTEM)
    if not isinstance(raw, dict):
        raise GeminiError("Gemini returned an unexpected structure.")

    vendor_in = raw.get("vendor") if isinstance(raw.get("vendor"), dict) else {}
    bank_in = raw.get("bank") if isinstance(raw.get("bank"), dict) else {}
    vendor = {k: _s(vendor_in.get(k)) for k in
              ("legal_name", "trade_name", "gstin", "pan", "address", "contact_person", "email", "phone", "category")}
    prods = vendor_in.get("products_supplied")
    vendor["products_supplied"] = [str(p).strip() for p in prods if str(p).strip()] if isinstance(prods, list) else []
    bank = {k: _s(bank_in.get(k)) for k in ("bank_name", "account_holder", "account_number", "ifsc")}
    for k in ("gstin", "pan"):
        if vendor[k]:
            vendor[k] = re.sub(r"\s+", "", vendor[k]).upper()
    if bank["ifsc"]:
        bank["ifsc"] = re.sub(r"\s+", "", bank["ifsc"]).upper()
    if bank["account_number"]:
        bank["account_number"] = re.sub(r"\s+", "", bank["account_number"])

    quotes = []
    for qd in raw.get("quotations") or []:
        if not isinstance(qd, dict) or not _s(qd.get("product_name")):
            continue
        days = to_number(qd.get("delivery_days"))
        quotes.append({
            "product_name": _s(qd.get("product_name")),
            "product_code": _s(qd.get("product_code")),
            "unit_price": to_number(qd.get("unit_price")),
            "quoted_quantity": to_number(qd.get("quoted_quantity")),
            "total_price": to_number(qd.get("total_price")),
            "taxes": _s(qd.get("taxes")),
            "delivery_days": int(days) if days is not None else None,
            "validity": _s(qd.get("validity")),
            "other_terms": _s(qd.get("other_terms")),
        })
    docs = []
    for d in raw.get("documents") or []:
        if isinstance(d, dict) and _s(d.get("file_name")):
            docs.append({"file_name": _s(d["file_name"]), "document_type": _s(d.get("document_type")) or "Other"})
    incons = [str(x).strip() for x in (raw.get("inconsistencies") or []) if str(x).strip()]
    qe = raw.get("quality_evidence") if isinstance(raw.get("quality_evidence"), dict) else {}
    certs = qe.get("certifications") if isinstance(qe.get("certifications"), list) else []
    quality = {"certifications": [str(c).strip() for c in certs if sc.clean(c)],
               "on_time": to_number(qe.get("on_time_delivery_percent")),
               "defect": to_number(qe.get("defect_or_return_rate_percent")),
               "other": _s(qe.get("other_evidence"))}
    return {"vendor": vendor, "bank": bank, "quotations": quotes, "documents": docs, "inconsistencies": incons,
            "quality": quality, "duplicate_matches": _parse_matches(raw.get("duplicate_matches"), existing),
            "summary": _s(raw.get("summary"))}


def _parse_matches(items, existing):
    names = {v["vendor_id"]: v["legal_name"] for v in existing or []}
    out = []
    for m in items or []:
        if not isinstance(m, dict):
            continue
        vid, sim = m.get("vendor_id"), to_number(m.get("similarity"))
        if vid in names and sim is not None and sim >= 60:
            reasons = [str(r) for r in (m.get("reasons") or []) if str(r).strip()][:3]
            out.append({"vendor_id": vid, "name": names[vid], "similarity": int(min(sim, 100)),
                        "reasons": reasons or ["Similar vendor details"]})
    return out


# ------------------------------------------------------ 3. vendor summary (fallback)
def fallback_summary(f):
    price = f.get("quoted_prices") or "no quoted price"
    deliv = f.get("delivery") or "no stated delivery time"
    return (f"{f.get('vendor')} is a {str(f.get('category') or 'supplier').lower()} offering {f.get('products') or 'no in-scope products'} "
            f"({price}; delivery: {deliv}). Documentation is {f.get('documentation', 'unknown').lower()} "
            f"and the duplicate risk is {str(f.get('duplicate_risk', 'unknown')).lower()}.")


# ----------------------------------------------------- 4. ranking 'Why?'
def explain_ranking(ctx):
    """Returns (text, source)."""
    try:
        text = generate(
            [{"text": "Application data (JSON):\n" + json.dumps(ctx, ensure_ascii=False, default=str) +
              f"\n\nExplain in 3-4 plain sentences why '{ctx['selected_vendor']}' is ranked #{ctx['selected_rank']}. "
              "Refer to the user's priorities (weights) and compare with at least one other vendor. "
              "Use ONLY numbers and facts from the data. Do not invent anything. Do not add headings or bullets."}],
            system="You explain vendor rankings for a procurement team. The scores were calculated by the application; "
                   "you only explain them.")
        return text.strip(), "ai"
    except GeminiError:
        return fallback_explanation(ctx), "fallback"


def fallback_explanation(ctx):
    w = ctx["weights_percent"]
    top_w = max(w, key=w.get)
    me = next(r for r in ctx["ranking"] if r["vendor"] == ctx["selected_vendor"])
    s = (f"{me['vendor']} is ranked #{ctx['selected_rank']} with a score of {me['score']}/100. "
         f"Your priorities place the most weight on {top_w} ({w[top_w]:g}%). "
         f"Its points are cost {me['cost_points']}, delivery {me['delivery_points']} and quality {me['quality_points']}.")
    others = [r for r in ctx["ranking"] if r["vendor"] != me["vendor"]]
    if others:
        o = others[0]
        s += f" For comparison, {o['vendor']} scores {o['score']}/100."
    return s


# ------------------------------------------------------------- 5. email
def build_email(f):
    """Ready-to-send purchase request email in the format given in the project brief."""
    subject = f"Purchase Request – {f['product_short']} – {f['quantity']}"
    price = f" at the quoted price of {f['unit_price']}" if f.get("unit_price") else ""
    lines = [f"Dear {f['vendor_team']} Team,", "",
             f"We would like to place a purchase request for {f['quantity']} of {f['product_short']}{price}.", "",
             "Order details:",
             f"• Product: {f['product']}",
             f"• Quantity: {f['quantity']}"]
    if f.get("unit_price"):
        lines.append(f"• Quoted unit price: {f['unit_price']}")
    if f.get("order_value"):
        lines.append(f"• Estimated order value: {f['order_value']}" +
                     (f" ({f['taxes']}, as per your quotation)" if f.get("taxes") else ""))
    lines.append(f"• Required delivery date: {f['delivery_date']}")
    if f.get("delivery_address"):
        lines.append(f"• Delivery address: {f['delivery_address']}")
    lines += ["", f"We request delivery by {f['delivery_date']}."]
    if f.get("special_instructions"):
        lines += ["", f"Special instructions: {f['special_instructions']}"]
    lines += ["", "Please confirm the availability of the required quantity and the delivery schedule.",
              "Kindly share your confirmation and invoice details.", "",
              "Regards,", "Procurement Team", COMPANY]
    return subject, "\n".join(lines)


EMAIL_SYSTEM = """You polish purchase-request emails. You receive a complete draft. Return ONLY JSON {"subject": "...", "body": "..."}.
Keep the SAME structure, line breaks, bullet points, greeting and signature. Keep every number, date, price and name EXACTLY.
Only improve grammar/flow of the sentences if needed. Do not add new information, discounts or promises. Keep it short."""


def generate_email(facts):
    """Returns (subject, body, source). The template guarantees a correct, ready-to-send format."""
    subject, body = build_email(facts)
    try:
        data = generate_json([{"text": json.dumps({"subject": subject, "body": body}, ensure_ascii=False)}],
                             system=EMAIL_SYSTEM)
        s2, b2 = str(data.get("subject", "")).strip(), str(data.get("body", "")).strip()
        must = [facts["quantity"], facts["delivery_date"], "Procurement Team"]
        if facts.get("unit_price"):
            must.append(facts["unit_price"])
        if s2 and b2 and all(m in b2 for m in must):
            return s2, b2, "ai"
    except GeminiError:
        pass
    return subject, body, "template"


# ----------------------------------------------------------- 6. assistant
ASSISTANT_SYSTEM = """You are the AI procurement assistant inside the M Beans & Bites procurement application.
Answer ONLY from the application data (JSON) provided in the message. The data comes from the live database
and from the application's own deterministic ranking calculation.
- If the answer is not in the data, say: "That information is not available in the application data."
- Never invent vendors, prices, dates, scores or order information.
- Use ₹ for prices. Keep answers short and clear (a few sentences or a short list).
- For ranking questions use the 'rankings' section; explain differences using the score breakdown given."""


def ask_assistant(question, context, history):
    convo = "\n".join(f"{m['role'].title()}: {m['content']}" for m in history[-6:])
    text = ("Application data (JSON):\n" + json.dumps(context, ensure_ascii=False, default=str) +
            (f"\n\nRecent conversation:\n{convo}" if convo else "") +
            f"\n\nUser question: {question}")
    return generate([{"text": text}], system=ASSISTANT_SYSTEM)
